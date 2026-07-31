from __future__ import annotations

from pathlib import Path
from typing import Any


class WorkbookReadError(ValueError):
    pass


def read_first_sheet(path: Path) -> list[list[Any]]:
    suffix = path.suffix.lower()
    try:
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(path, read_only=True, data_only=True)
            sheet = workbook[workbook.sheetnames[0]]
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        if suffix == ".xls":
            import xlrd

            workbook = xlrd.open_workbook(path)
            sheet = workbook.sheet_by_index(0)
            return [sheet.row_values(i) for i in range(sheet.nrows)]
    except Exception as exc:  # library exceptions vary by file format
        raise WorkbookReadError("فایل Excel قابل خواندن نیست یا آسیب دیده است") from exc
    raise WorkbookReadError("فقط فایل‌های xlsx و xls قابل قبول هستند")
